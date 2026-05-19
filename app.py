import streamlit as st
from datetime import date, datetime
from zoneinfo import ZoneInfo
from io import BytesIO
import json
import requests
from pypdf import PdfReader
from openai import OpenAI
from html import escape
import re

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

import base64
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import gspread
from google.oauth2.service_account import Credentials

# =========================================================
# CONFIGURAÇÃO
# =========================================================
st.set_page_config(
    page_title="LexCalx - Cálculos jurídicos",
    page_icon="⚖️",
    layout="wide",
)

APP_DIR = Path(__file__).resolve().parent
CSS_PATH = APP_DIR / "style.css"
LOGO_PATH = APP_DIR / "ralogo.png"
HISTORICO_CALCULOS_PATH = APP_DIR / "historico_calculos.json"


def imagem_base64(caminho: Path) -> str:
    if not caminho.exists():
        st.error(f"Imagem não encontrada: {caminho}")
        return ""

    return base64.b64encode(caminho.read_bytes()).decode()

def obter_data_calculo():
    return datetime.now(ZoneInfo("America/Sao_Paulo")).date()

def registrar_solicitacao_calculo(usuario_logado, gcpj, numero_processo, data_calculo):
    url = st.secrets.get("LEXCALX_HISTORICO_URL", "")
    token = st.secrets.get("LEXCALX_HISTORICO_TOKEN", "")

    if not url or not token:
        st.error("Configuração do histórico não encontrada nos secrets.")
        st.stop()

    payload = {
        "token": token,
        "usuario": usuario_logado,
        "gcpj": gcpj,
        "numero_processo": numero_processo,
        "data_calculo": data_calculo,
        "modelo_openai": st.secrets.get("OPENAI_MODEL", "gpt-4.1"),
    }

    try:
        resposta = requests.post(url, json=payload, timeout=20)
        resposta.raise_for_status()

        retorno = resposta.json()

        if not retorno.get("ok"):
            st.error(f"Erro ao registrar histórico: {retorno.get('erro', 'Erro desconhecido')}")
            st.stop()

    except Exception as erro:
        st.error(f"Erro ao registrar histórico do cálculo: {erro}")
        st.stop()

def carregar_css():
    if not CSS_PATH.exists():
        st.error(f"CSS não encontrado: {CSS_PATH}")
        return

    css = CSS_PATH.read_text(encoding="utf-8")
    logo_bg = imagem_base64(LOGO_PATH)

    css = css.replace("{{LOGO_BG}}", logo_bg)

    st.markdown(
        f"<style>{css}</style>",
        unsafe_allow_html=True,
    )

carregar_css()

def imagem_base64(caminho: Path) -> str:
    if not caminho.exists():
        return ""
    return base64.b64encode(caminho.read_bytes()).decode()


LOGO_BG = imagem_base64(LOGO_PATH)

# =========================================================
# FUNÇÕES AUXILIARES
# =========================================================
def conectar_google_sheets():
    escopos = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]

    credenciais = Credentials.from_service_account_info(
        st.secrets["gcp_service_account"],
        scopes=escopos,
    )

    cliente = gspread.authorize(credenciais)

    sheet_id = st.secrets["GOOGLE_SHEET_ID"]

    return cliente.open_by_key(sheet_id)

def formatar_data_br(data_valor):
    if not data_valor:
        return ""
    return data_valor.strftime("%d/%m/%Y")

def converter_valor_br_para_float(valor_texto):
    if not valor_texto:
        return None

    try:
        valor_limpo = str(valor_texto).strip()
        valor_limpo = valor_limpo.replace("R$", "").strip()
        valor_limpo = valor_limpo.replace(".", "")
        valor_limpo = valor_limpo.replace(",", ".")

        return float(valor_limpo)
    except Exception:
        return None

def formatar_valor_texto_br(valor_texto):
    valor_float = converter_valor_br_para_float(valor_texto)

    if valor_float is None:
        return ""

    return f"{valor_float:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def normalizar_campo_valor_dano_material(indice):
    chave = f"valor_dano_material_{indice}"

    valor_digitado = st.session_state.get(chave, "")

    if valor_digitado:
        st.session_state[chave] = formatar_valor_texto_br(valor_digitado)

def formatar_valor_br(valor):
    if valor in (None, ""):
        return ""

    try:
        valor = float(valor)
        return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except Exception:
        return str(valor)

def extrair_texto_pdf(arquivo_pdf) -> str:
    if arquivo_pdf is None:
        return ""

    try:
        leitor = PdfReader(BytesIO(arquivo_pdf.read()))
        textos = []

        for pagina in leitor.pages:
            texto = pagina.extract_text() or ""
            textos.append(texto)

        return "\n\n".join(textos).strip()

    except Exception as erro:
        return f"ERRO AO LER PDF: {erro}"

def separar_resultado_lexcalx(texto_resultado):
    texto_resultado = texto_resultado or ""

    inicio_resumo = "[INICIO_RESUMO_PROCESSO]"
    fim_resumo = "[FIM_RESUMO_PROCESSO]"

    inicio_historico = "[INICIO_HISTORICO_COMPLETO]"
    fim_historico = "[FIM_HISTORICO_COMPLETO]"

    def extrair_trecho(texto, inicio, fim):
        if inicio not in texto or fim not in texto:
            return ""

        return texto.split(inicio, 1)[1].split(fim, 1)[0].strip()

    resumo = extrair_trecho(texto_resultado, inicio_resumo, fim_resumo)
    historico = extrair_trecho(texto_resultado, inicio_historico, fim_historico)

    if not resumo:
        resumo = texto_resultado.strip()

    if not historico:
        historico = texto_resultado.strip()

    return resumo, historico

def montar_prompt(dados: dict, texto_decisao: str) -> str:
    return f"""

Você é o LexCalx, atuando como contador judicial profissional sênior, especializado em cálculos jurídicos cíveis brasileiros.

Sua tarefa é analisar a decisão judicial enviada pelo usuário, identificar todas as condenações existentes e elaborar a memória de cálculo até a DATA DO CÁLCULO informada.

Você deve atuar com rigor técnico-contábil, considerando que o cálculo será apresentado em processo judicial.

FONTES E ÍNDICES:
- Use somente dados oficiais do Governo Brasileiro.
- Para SELIC, utilize somente dados oficiais do Banco Central do Brasil.
- Para IPCA, INPC ou índices de inflação, utilize somente dados oficiais do IBGE ou dados oficiais reproduzidos pelo Banco Central do Brasil.
- Para tabelas judiciais, utilize somente fonte oficial do respectivo Tribunal.
- Não use blogs, sites privados, calculadoras particulares, artigos, notícias, JusBrasil, Escavador, Migalhas ou fontes não oficiais.
- Não estime índice, não use média aproximada e não invente percentual.
- Se o índice do mês final ainda não estiver disponível, use o último índice oficial disponível e informe isso de forma objetiva.
    
ATENÇÃO:
Você deve calcular qualquer tipo de decisão, não apenas dano moral.
A decisão pode conter dano moral, dano material, repetição de indébito, multa, astreintes, honorários, obrigação de pagar, valores parcelados, condenação solidária, improcedência parcial ou condenação ilíquida.

REGRA PRINCIPAL:
- Use o DISPOSITIVO da decisão como fonte principal.
- Não use pedido inicial como condenação, salvo se a decisão expressamente acolher esse valor.
- Não trate como "dado faltante" uma verba que NÃO foi objeto de condenação.
- Se a decisão não condenou em dano material, escreva: "Não houve condenação em dano material."
- Se a decisão não fixou honorários, escreva: "Não houve condenação em honorários nesta fase."
- Somente aponte pendência quando existir condenação e faltar dado essencial para liquidar.
- Custas não devem ser calculadas, nem incluídas no resumo, nem incluídas no total.
- Se a decisão mencionar custas, ignore para fins de cálculo.

DATA DO CÁLCULO:
- A atualização deve ser feita até: {dados["data_calculo"]}.
- Essa é a data limite do cálculo informada no campo "Atualizar o cálculo até".
- Por padrão, esse campo vem preenchido com a data do dia da execução.
- Se o usuário alterar essa data, use a data alterada como limite do cálculo.
- Se algum índice oficial necessário não estiver disponível até essa data, use o último índice oficial disponível e informe isso expressamente.

RESULTADO NUMÉRICO OBRIGATÓRIO:
- A versão resumida e o histórico completo devem apresentar valores monetários finais calculados numericamente.
- É proibido usar variáveis, letras, símbolos ou placeholders no lugar de valores.
- É proibido usar expressões como: [A], [B], [C], R$ [A], R$ [B], "a calcular", "ver cálculo abaixo", "consultar histórico" ou semelhantes.
- Se houver condenação líquida e critérios de atualização definidos, calcule o valor atualizado e apresente o resultado em R$.
- Se o índice do mês final ainda não estiver disponível, use o último índice oficial disponível e apresente o valor numérico calculado com esse último índice.
- Se alguma verba realmente não puder ser calculada por falta de dado essencial da decisão, não coloque placeholder na tabela; informe essa verba em "Pendências".

ART. 523 DO CPC:
- Opção selecionada para art. 523: {dados["inserir_art_523"]}.
- Este campo é uma ordem operacional do usuário.
- NÃO analise cabimento jurídico do art. 523.
- NÃO diga que o art. 523 não se aplica por ser Juizado Especial, sentença não definitiva, fase processual inadequada ou qualquer outro motivo.
- Aplique exatamente a opção selecionada pelo usuário.

Regras obrigatórias:
- Se a opção for "Não", não acrescente multa nem honorários do art. 523.
- Se a opção for "Multa 10%", após apurar o total atualizado da condenação, acrescente somente multa de 10% sobre o total atualizado.
- Se a opção for "Honorários 10%", após apurar o total atualizado da condenação, acrescente somente honorários de 10% sobre o total atualizado.
- Se a opção for "Ambos", após apurar o total atualizado da condenação, acrescente multa de 10% e honorários de 10%, totalizando acréscimo de 20% sobre o total atualizado.

Fórmulas obrigatórias:
- Multa 10% = total atualizado x 10%.
- Honorários 10% = total atualizado x 10%.
- Ambos = total atualizado x 20%.
- Total geral com apenas multa = total atualizado x 1,10.
- Total geral com apenas honorários = total atualizado x 1,10.
- Total geral com ambos = total atualizado x 1,20.

Não confunda honorários do art. 523 com honorários sucumbenciais fixados na decisão.

COMO PROCEDER:
1. Leia a decisão inteira.
2. Identifique o processo, partes, data da decisão e data da homologação/publicação, se constarem.
3. Localize o dispositivo.
4. Identifique cada verba condenatória.
5. Para cada verba, identifique:
   - natureza da verba;
   - valor principal;
   - termo inicial dos juros;
   - termo inicial da correção monetária;
   - índice de juros;
   - índice de correção monetária;
   - percentual de honorários, se houver;
   - multa, se houver;

6. Use as datas informadas pelo usuário quando forem compatíveis com a decisão.
7. Se a decisão trouxer data mais específica que o usuário não preencheu, use a data constante da decisão.
8. Se houver divergência entre o dado do usuário e a decisão, aponte a divergência e explique qual dado foi usado.
9. Realize o cálculo matemático até a data do cálculo.
10. Apresente memória de cálculo clara, objetiva e auditável.

REGRAS SOBRE ÍNDICES:
- Se a decisão determinar SELIC, IPCA, INPC, IGP-M, TJMG, tabela judicial, poupança, juros de 1% ao mês ou outro critério, aplique exatamente o critério determinado.
- Se precisar consultar índice atualizado, use exclusivamente fonte oficial do Governo Brasileiro ou do Tribunal competente.
- Para SELIC, use Banco Central do Brasil.
- Para IPCA ou INPC, use IBGE ou Banco Central do Brasil quando a série oficial estiver disponível.
- Para tabela judicial, use o site oficial do respectivo Tribunal.
- Não use fonte privada, calculadora privada, blog, artigo, notícia ou site jurídico não oficial.
- Não invente índice, data, valor, taxa ou percentual.
- Não use estimativa, média aproximada ou projeção.
- Se o índice do mês final ainda não estiver disponível, use o último índice oficial disponível e informe isso de forma objetiva.
- Quando usar índice pesquisado, informe a fonte oficial utilizada.
- Não misture critérios diferentes se a decisão for clara.

REGRAS SOBRE JUROS E CORREÇÃO:
- Juros e correção devem ser calculados separadamente quando a decisão assim permitir.
- Quando a decisão determinar SELIC como índice único, não some correção monetária separada, salvo se a decisão expressamente mandar.
- Quando a decisão determinar SELIC deduzida de IPCA e IPCA separado, explique a metodologia.
- Se houver juros desde o evento danoso, use a data do evento danoso informada ou identificada na decisão.
- Se houver correção desde o arbitramento, use a data da sentença ou da decisão que arbitrou o valor.

REGRAS SOBRE VERBAS NÃO EXISTENTES:
- Não escreva "faltam dados" para dano material, honorários ou multa se a decisão não condenou nessas verbas.
- Liste essas verbas apenas no resumo como "não houve condenação", quando necessário.

FORMATO DA RESPOSTA:

Você deve devolver obrigatoriamente DUAS versões:

1. Uma versão resumida, própria para juntada no processo.
2. Um histórico completo, para conferência interna.

Use exatamente os marcadores abaixo.

[INICIO_RESUMO_PROCESSO]

# CÁLCULO DA DECISÃO

Cálculo da decisão proferida no processo {dados["numero_processo"]}, conforme dispositivo da sentença/decisão.

Data de atualização do cálculo: {dados["data_calculo"]}.

## Resumo do cálculo

Monte uma tabela objetiva com as colunas:

Verba | Valor principal | Critério de atualização | Valor atualizado | Art. 523 | Total

## Índices utilizados

Explique de forma sucinta, em poucas linhas:
- qual índice de correção monetária foi usado;
- qual índice ou taxa de juros foi usado;
- qual termo inicial foi considerado;
- se foi usado último índice oficial disponível, informe isso de forma objetiva.

## Total

Informe o total atualizado da condenação.

Se houver aplicação do art. 523, informe separadamente:
- multa;
- honorários;
- total geral.

Não inclua histórico longo, metodologia extensa, links grandes, jurisprudência, transcrição da decisão ou explicações desnecessárias.

[FIM_RESUMO_PROCESSO]

[INICIO_HISTORICO_COMPLETO]

# MEMÓRIA COMPLETA DE CÁLCULO - LEXCALX

## 1. Dados identificados na decisão
- GCPJ: {dados["gcpj"]}
- Processo:
- Autor:
- Réu:
- Data da decisão:
- Data do cálculo:

## 2. Resumo das condenações
Monte uma tabela com:
Verba | Valor principal | Termo inicial juros | Termo inicial correção | Índice/Critério | Resultado atualizado | Observação

## 3. Cálculo detalhado por verba
Para cada verba condenatória, apresente:
- valor principal;
- índice aplicado;
- período utilizado;
- fórmula;
- valor de correção;
- valor de juros;
- valor atualizado.

## 4. Honorários, custas e multa
Informe apenas se houver condenação.
Se não houver, informe expressamente que não houve condenação.

## 5. Total atualizado
Apresente o total geral atualizado até a data do cálculo.

## 6. Pendências
Liste somente pendências reais que impeçam o cálculo de verba efetivamente condenada.

[FIM_HISTORICO_COMPLETO]

DADOS INFORMADOS PELO USUÁRIO:

GCPJ: {dados["gcpj"]}
Número do processo: {dados["numero_processo"]}

Data do cálculo: {dados["data_calculo"]}
Opção art. 523: {dados["inserir_art_523"]}
Multa art. 523: {dados["percentual_multa_art_523"]}
Honorários art. 523: {dados["percentual_honorarios_art_523"]}
Acréscimo total art. 523: {dados["percentual_total_art_523"]}

Data da distribuição: {dados["data_distribuicao"]}
Data da citação: {dados["data_citacao"]}
Data do evento danoso: {dados["data_evento_danoso"]}

Data juros dano moral: {dados["data_juros_dano_moral"]}
Data correção monetária dano moral: {dados["data_correcao_dano_moral"]}

Data juros dano material: {dados["data_juros_dano_material"]}
Data correção monetária dano material: {dados["data_correcao_dano_material"]}

Data da decisão: {dados["data_decisao"]}
Data da publicação: {dados["data_publicacao"]}

Danos materiais informados pelo usuário:
{dados["danos_materiais"]}

TEXTO INTEGRAL DA DECISÃO:

{texto_decisao}
"""

def gerar_planilha_calculo(dados, resultado_gpt):
    wb = Workbook()

    ws_dados = wb.active
    ws_dados.title = "Dados informados"

    ws_resultados = wb.create_sheet("Resultado")
    ws_danos = wb.create_sheet("Danos materiais")

    cor_titulo = "1F3A4D"
    cor_header = "D7E2EB"
    cor_texto_claro = "FFFFFF"

    fonte_titulo = Font(bold=True, size=14, color=cor_texto_claro)
    fonte_header = Font(bold=True, size=11, color="000000")
    fonte_padrao = Font(size=10, color="000000")

    fill_titulo = PatternFill("solid", fgColor=cor_titulo)
    fill_header = PatternFill("solid", fgColor=cor_header)

    borda = Border(
        left=Side(style="thin", color="AFC0CC"),
        right=Side(style="thin", color="AFC0CC"),
        top=Side(style="thin", color="AFC0CC"),
        bottom=Side(style="thin", color="AFC0CC"),
    )

    alinhamento_padrao = Alignment(
        vertical="top",
        horizontal="left",
        wrap_text=True,
    )

    # =====================================================
    # ABA 1 - DADOS INFORMADOS
    # =====================================================
    ws_dados.merge_cells("A1:B1")
    ws_dados["A1"] = "LexCalx - Dados informados"
    ws_dados["A1"].font = fonte_titulo
    ws_dados["A1"].fill = fill_titulo
    ws_dados["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws_dados["A3"] = "Campo"
    ws_dados["B3"] = "Informação"

    for celula in ws_dados[3]:
        celula.font = fonte_header
        celula.fill = fill_header
        celula.border = borda
        celula.alignment = alinhamento_padrao

    dados_linhas = [
        ("GCPJ", dados.get("gcpj", "")),
        ("Número do processo", dados.get("numero_processo", "")),
        ("Data do cálculo", dados.get("data_calculo", "")),
        ("Opção art. 523", dados.get("inserir_art_523", "")),
        ("Multa art. 523", dados.get("percentual_multa_art_523", "")),
        ("Honorários art. 523", dados.get("percentual_honorarios_art_523", "")),
        ("Acréscimo total art. 523", dados.get("percentual_total_art_523", "")),
        ("Data da distribuição", dados.get("data_distribuicao", "")),
        ("Data da citação", dados.get("data_citacao", "")),
        ("Data do evento danoso", dados.get("data_evento_danoso", "")),
        ("Data juros dano moral", dados.get("data_juros_dano_moral", "")),
        ("Data correção monetária dano moral", dados.get("data_correcao_dano_moral", "")),
        ("Data juros dano material", dados.get("data_juros_dano_material", "")),
        ("Data correção monetária dano material", dados.get("data_correcao_dano_material", "")),
        ("Data da decisão", dados.get("data_decisao", "")),
        ("Data da publicação", dados.get("data_publicacao", "")),
    ]

    for linha_idx, linha in enumerate(dados_linhas, start=4):
        ws_dados.cell(row=linha_idx, column=1, value=linha[0])
        ws_dados.cell(row=linha_idx, column=2, value=linha[1])

        for col_idx in range(1, 3):
            celula = ws_dados.cell(row=linha_idx, column=col_idx)
            celula.font = fonte_padrao
            celula.border = borda
            celula.alignment = alinhamento_padrao

    ws_dados.column_dimensions["A"].width = 38
    ws_dados.column_dimensions["B"].width = 45
    ws_dados.freeze_panes = "A4"

    # =====================================================
    # ABA 2 - DANOS MATERIAIS
    # =====================================================
    ws_danos.merge_cells("A1:C1")
    ws_danos["A1"] = "LexCalx - Danos materiais informados"
    ws_danos["A1"].font = fonte_titulo
    ws_danos["A1"].fill = fill_titulo
    ws_danos["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers_danos = ["Valor", "Parcelas", "Data inicial"]

    for col_idx, header in enumerate(headers_danos, start=1):
        celula = ws_danos.cell(row=3, column=col_idx, value=header)
        celula.font = fonte_header
        celula.fill = fill_header
        celula.border = borda
        celula.alignment = alinhamento_padrao

    danos_texto = dados.get("danos_materiais", "Não informado.")

    if danos_texto and danos_texto != "Não informado.":
        linhas_danos = danos_texto.split("\n")

        for row_idx, linha in enumerate(linhas_danos, start=4):
            partes = linha.split("|")

            valor = partes[0].replace(f"{row_idx - 3}. Valor:", "").strip() if len(partes) > 0 else ""
            parcelas = partes[1].replace("Parcelas:", "").strip() if len(partes) > 1 else ""
            data_inicial = partes[2].replace("Data inicial:", "").strip() if len(partes) > 2 else ""

            ws_danos.cell(row=row_idx, column=1, value=valor)
            ws_danos.cell(row=row_idx, column=2, value=parcelas)
            ws_danos.cell(row=row_idx, column=3, value=data_inicial)

            for col_idx in range(1, 4):
                celula = ws_danos.cell(row=row_idx, column=col_idx)
                celula.font = fonte_padrao
                celula.border = borda
                celula.alignment = alinhamento_padrao
    else:
        ws_danos.cell(row=4, column=1, value="Não informado.")
        ws_danos.cell(row=4, column=1).font = fonte_padrao

    ws_danos.column_dimensions["A"].width = 25
    ws_danos.column_dimensions["B"].width = 15
    ws_danos.column_dimensions["C"].width = 20
    ws_danos.freeze_panes = "A4"

    # =====================================================
    # ABA 3 - RESULTADO CHATGPT
    # =====================================================
    ws_resultados.merge_cells("A1:B1")
    ws_resultados["A1"] = "LexCalx - Resultado do cálculo"
    ws_resultados["A1"].font = fonte_titulo
    ws_resultados["A1"].fill = fill_titulo
    ws_resultados["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws_resultados["A3"] = "Fonte"
    ws_resultados["B3"] = "Resultado"

    for celula in ws_resultados[3]:
        celula.font = fonte_header
        celula.fill = fill_header
        celula.border = borda
        celula.alignment = alinhamento_padrao

    ws_resultados.cell(row=4, column=1, value="ChatGPT")
    ws_resultados.cell(row=4, column=2, value=resultado_gpt)

    for col_idx in range(1, 3):
        celula = ws_resultados.cell(row=4, column=col_idx)
        celula.font = fonte_padrao
        celula.border = borda
        celula.alignment = alinhamento_padrao

    ws_resultados.row_dimensions[4].height = 240
    ws_resultados.column_dimensions["A"].width = 18
    ws_resultados.column_dimensions["B"].width = 110
    ws_resultados.freeze_panes = "A4"

    # =====================================================
    # AJUSTES GERAIS
    # =====================================================
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = alinhamento_padrao

    arquivo_excel = BytesIO()
    wb.save(arquivo_excel)
    arquivo_excel.seek(0)

    return arquivo_excel

def texto_pdf(texto):
    if texto in (None, ""):
        return ""

    texto = str(texto)
    texto = escape(texto)
    texto = texto.replace("\n", "<br/>")

    return texto

def limpar_markdown_basico(texto):
    texto = escape(str(texto or ""))

    texto = re.sub(
        r"\*\*(.*?)\*\*",
        r"<b>\1</b>",
        texto,
    )

    return texto

def linha_tabela_markdown(linha):
    linha = linha.strip()
    return linha.startswith("|") and linha.endswith("|")

def linha_separadora_tabela(linha):
    conteudo = linha.replace("|", "").replace("-", "").replace(":", "").strip()
    return conteudo == ""

def quebrar_linha_tabela(linha):
    return [parte.strip() for parte in linha.strip().strip("|").split("|")]

def markdown_resumo_para_html(texto):
    linhas = str(texto or "").splitlines()
    html = []
    tabela = []

    def fechar_tabela():
        nonlocal tabela

        if not tabela:
            return

        html.append('<table class="resumo-tabela">')

        for idx, linha_tabela in enumerate(tabela):
            tag = "th" if idx == 0 else "td"
            html.append("<tr>")

            for celula in linha_tabela:
                html.append(f"<{tag}>{limpar_markdown_basico(celula)}</{tag}>")

            html.append("</tr>")

        html.append("</table>")
        tabela = []

    for linha in linhas:
        linha_limpa = linha.strip()

        if not linha_limpa:
            fechar_tabela()
            continue

        if linha_tabela_markdown(linha_limpa):
            if linha_separadora_tabela(linha_limpa):
                continue

            tabela.append(quebrar_linha_tabela(linha_limpa))
            continue

        fechar_tabela()

        if linha_limpa.startswith("# "):
            html.append(f"<h1>{limpar_markdown_basico(linha_limpa.replace('# ', '', 1))}</h1>")
        elif linha_limpa.startswith("## "):
            html.append(f"<h2>{limpar_markdown_basico(linha_limpa.replace('## ', '', 1))}</h2>")
        elif linha_limpa.startswith("### "):
            html.append(f"<h3>{limpar_markdown_basico(linha_limpa.replace('### ', '', 1))}</h3>")
        elif linha_limpa.startswith("- "):
            html.append(f"<p>• {limpar_markdown_basico(linha_limpa.replace('- ', '', 1))}</p>")
        else:
            html.append(f"<p>{limpar_markdown_basico(linha_limpa)}</p>")

    fechar_tabela()

    return "\n".join(html)

def gerar_pdf_calculo(dados, resultado_gpt):
    arquivo_pdf = BytesIO()

    doc = SimpleDocTemplate(
        arquivo_pdf,
        pagesize=A4,
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    estilos = getSampleStyleSheet()

    estilo_titulo = ParagraphStyle(
        "TituloLexCalx",
        parent=estilos["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        textColor=colors.HexColor("#1F3A4D"),
        alignment=1,
        spaceAfter=18,
    )

    estilo_subtitulo = ParagraphStyle(
        "SubtituloLexCalx",
        parent=estilos["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        textColor=colors.HexColor("#1F3A4D"),
        spaceBefore=10,
        spaceAfter=8,
    )

    estilo_normal = ParagraphStyle(
        "TextoLexCalx",
        parent=estilos["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.black,
    )

    estilo_resultado = ParagraphStyle(
        "ResultadoLexCalx",
        parent=estilos["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.black,
        spaceAfter=8,
    )

    elementos = []

    elementos.append(Paragraph("LexCalx - Cálculo Jurídico", estilo_titulo))

    # =====================================================
    # DADOS INFORMADOS
    # =====================================================
    elementos.append(Paragraph("Dados informados", estilo_subtitulo))

    dados_linhas = [
        ["Campo", "Informação"],
         ["GCPJ", dados.get("gcpj", "")],
        ["Número do processo", dados.get("numero_processo", "")],
        ["Data do cálculo", dados.get("data_calculo", "")],
        ["Opção art. 523", dados.get("inserir_art_523", "")],
        ["Multa art. 523", dados.get("percentual_multa_art_523", "")],
        ["Honorários art. 523", dados.get("percentual_honorarios_art_523", "")],
        ["Acréscimo total art. 523", dados.get("percentual_total_art_523", "")],
        ["Data da distribuição", dados.get("data_distribuicao", "")],
        ["Data da citação", dados.get("data_citacao", "")],
        ["Data do evento danoso", dados.get("data_evento_danoso", "")],
        ["Data juros dano moral", dados.get("data_juros_dano_moral", "")],
        ["Data correção monetária dano moral", dados.get("data_correcao_dano_moral", "")],
        ["Data juros dano material", dados.get("data_juros_dano_material", "")],
        ["Data correção monetária dano material", dados.get("data_correcao_dano_material", "")],
        ["Data da decisão", dados.get("data_decisao", "")],
        ["Data da publicação", dados.get("data_publicacao", "")],
    ]

    tabela_dados = Table(
        dados_linhas,
        colWidths=[7 * cm, 10 * cm],
        repeatRows=1,
    )

    tabela_dados.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D7E2EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AFC0CC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )

    elementos.append(tabela_dados)
    elementos.append(Spacer(1, 12))

    # =====================================================
    # DANOS MATERIAIS
    # =====================================================
    elementos.append(Paragraph("Danos materiais informados", estilo_subtitulo))

    danos_texto = dados.get("danos_materiais", "Não informado.")

    if danos_texto and danos_texto != "Não informado.":
        linhas_danos = [["Valor", "Parcelas", "Data inicial"]]

        for idx, linha in enumerate(danos_texto.split("\n"), start=1):
            partes = linha.split("|")

            valor = partes[0].replace(f"{idx}. Valor:", "").strip() if len(partes) > 0 else ""
            parcelas = partes[1].replace("Parcelas:", "").strip() if len(partes) > 1 else ""
            data_inicial = partes[2].replace("Data inicial:", "").strip() if len(partes) > 2 else ""

            linhas_danos.append([valor, parcelas, data_inicial])

        tabela_danos = Table(
            linhas_danos,
            colWidths=[6 * cm, 4 * cm, 5 * cm],
            repeatRows=1,
        )

        tabela_danos.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D7E2EB")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AFC0CC")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )

        elementos.append(tabela_danos)

    else:
        elementos.append(Paragraph("Não informado.", estilo_normal))

    elementos.append(Spacer(1, 12))

    # =====================================================
    # RESULTADO CHATGPT
    # =====================================================
    elementos.append(Paragraph("Resultado do cálculo - LexCalx", estilo_subtitulo))

    resultado_formatado = texto_pdf(resultado_gpt)

    elementos.append(
        Paragraph(
            resultado_formatado,
            estilo_resultado,
        )
    )

    doc.build(elementos)

    arquivo_pdf.seek(0)

    return arquivo_pdf

def gerar_pdf_resumo_processo(dados, resumo_processo):
    arquivo_pdf = BytesIO()

    doc = SimpleDocTemplate(
        arquivo_pdf,
        pagesize=landscape(A4),
        rightMargin=1.2 * cm,
        leftMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
    )

    estilos = getSampleStyleSheet()

    estilo_titulo = ParagraphStyle(
        "TituloResumoLexCalx",
        parent=estilos["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        textColor=colors.HexColor("#1F3A4D"),
        alignment=1,
        spaceAfter=14,
    )

    estilo_h2 = ParagraphStyle(
        "H2ResumoLexCalx",
        parent=estilos["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        textColor=colors.HexColor("#1F3A4D"),
        spaceBefore=8,
        spaceAfter=6,
    )

    estilo_normal = ParagraphStyle(
        "TextoResumoLexCalx",
        parent=estilos["BodyText"],
        fontName="Helvetica",
        fontSize=8.8,
        leading=11,
        textColor=colors.black,
        spaceAfter=6,
    )

    estilo_celula = ParagraphStyle(
        "CelulaResumoLexCalx",
        parent=estilos["BodyText"],
        fontName="Helvetica",
        fontSize=7.6,
        leading=9.2,
        textColor=colors.black,
    )

    estilo_cabecalho = ParagraphStyle(
        "CabecalhoResumoLexCalx",
        parent=estilos["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=7.8,
        leading=9.5,
        textColor=colors.black,
    )

    elementos = []
    elementos.append(Paragraph("Cálculo da Decisão", estilo_titulo))

    linhas = str(resumo_processo or "").splitlines()
    tabela = []

    largura_util = 27.0 * cm

    def adicionar_tabela():
        nonlocal tabela

        if not tabela:
            return

        qtd_colunas = len(tabela[0])

        dados_tabela = []

        for idx_linha, linha_tabela in enumerate(tabela):
            estilo = estilo_cabecalho if idx_linha == 0 else estilo_celula

            dados_tabela.append(
                [
                    Paragraph(
                        limpar_markdown_basico(celula),
                        estilo,
                    )
                    for celula in linha_tabela
                ]
            )

        if qtd_colunas == 6:
            larguras = [
                3.0 * cm,
                3.0 * cm,
                8.5 * cm,
                3.2 * cm,
                4.3 * cm,
                5.0 * cm,
            ]
        else:
            larguras = [largura_util / qtd_colunas] * qtd_colunas

        tabela_pdf = Table(
            dados_tabela,
            colWidths=larguras,
            repeatRows=1,
        )

        tabela_pdf.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D7E2EB")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#AFC0CC")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )

        elementos.append(tabela_pdf)
        elementos.append(Spacer(1, 8))

        tabela = []

    for linha in linhas:
        linha_limpa = linha.strip()

        if not linha_limpa:
            adicionar_tabela()
            continue

        if linha_tabela_markdown(linha_limpa):
            if linha_separadora_tabela(linha_limpa):
                continue

            tabela.append(quebrar_linha_tabela(linha_limpa))
            continue

        adicionar_tabela()

        if linha_limpa.startswith("# "):
            elementos.append(
                Paragraph(
                    limpar_markdown_basico(linha_limpa.replace("# ", "", 1)),
                    estilo_h2,
                )
            )
        elif linha_limpa.startswith("## "):
            elementos.append(
                Paragraph(
                    limpar_markdown_basico(linha_limpa.replace("## ", "", 1)),
                    estilo_h2,
                )
            )
        elif linha_limpa.startswith("### "):
            elementos.append(
                Paragraph(
                    limpar_markdown_basico(linha_limpa.replace("### ", "", 1)),
                    estilo_h2,
                )
            )
        elif linha_limpa.startswith("- "):
            elementos.append(
                Paragraph(
                    "• " + limpar_markdown_basico(linha_limpa.replace("- ", "", 1)),
                    estilo_normal,
                )
            )
        else:
            elementos.append(
                Paragraph(
                    limpar_markdown_basico(linha_limpa),
                    estilo_normal,
                )
            )

    adicionar_tabela()

    doc.build(elementos)

    arquivo_pdf.seek(0)

    return arquivo_pdf

def chamar_gpt(prompt: str) -> str:
    try:
        api_key = st.secrets.get("OPENAI_API_KEY", "")
        modelo = st.secrets.get("OPENAI_MODEL", "gpt-4.1")

        if not api_key:
            return (
                "Chave OPENAI_API_KEY não configurada. "
                "Inclua OPENAI_API_KEY no secrets.toml ou nos Secrets do Streamlit Cloud."
            )

        client = OpenAI(api_key=api_key)

        resposta = client.responses.create(
            model=modelo,
            temperature=0,
            instructions="""
Você é o LexCalx, atuando como contador judicial profissional sênior brasileiro.

Você deve analisar decisões judiciais cíveis e elaborar memória de cálculo objetiva, com rigor técnico-contábil, utilizando somente dados oficiais do Governo Brasileiro.

REGRAS OBRIGATÓRIAS:
- Calcule somente verbas efetivamente condenadas no dispositivo da decisão.
- Use o dispositivo da decisão como fonte principal.
- Não considere pedido inicial como condenação, salvo acolhimento expresso no dispositivo.
- Não diga que falta dado de verba inexistente.
- Não calcule custas e não inclua custas no total.
- Atualize os valores até a data do cálculo informada pelo usuário.
- Não use estimativa, média aproximada, projeção ou arredondamento livre.
- Não altere a metodologia entre execuções.
- Ao usar índice oficial, informe objetivamente qual índice foi usado e até qual competência/data.
- Se o índice do mês final ainda não estiver disponível, use o último índice oficial disponível e informe isso objetivamente.
- Não invente índices, datas, valores, taxas ou percentuais.
- Não use blogs, sites privados, calculadoras particulares, artigos, notícias ou sites jurídicos não oficiais.
- Use exclusivamente fontes oficiais do Governo Brasileiro ou do Tribunal competente.
- Para SELIC, use somente Banco Central do Brasil.
- Para IPCA ou INPC, use somente IBGE ou Banco Central do Brasil quando a série oficial estiver disponível.
- Para tabela judicial, use somente site oficial do respectivo Tribunal.

ART. 523:
- A opção de art. 523 é comando operacional do usuário.
- Se a opção for "Não", não aplique art. 523.
- Se a opção for "Multa 10%", aplique somente 10% de multa sobre o total atualizado.
- Se a opção for "Honorários 10%", aplique somente 10% de honorários sobre o total atualizado.
- Se a opção for "Ambos", aplique 10% de multa e 10% de honorários, totalizando 20% sobre o total atualizado.
- Não analise cabimento jurídico do art. 523.
- Não recuse aplicação do art. 523 por rito, fase processual, Juizado Especial ou qualquer outro fundamento.

RESULTADO:
- A versão resumida e o histórico completo devem trazer valores numéricos calculados em R$.
- É proibido usar variáveis, letras, símbolos ou placeholders no lugar dos valores.
- É proibido usar expressões como [A], [B], [C], R$ [A], R$ [B], "a calcular", "ver cálculo abaixo", "consultar histórico" ou semelhantes.
- Se não for possível calcular alguma verba condenada por falta de dado essencial, explique exatamente a pendência, sem colocar placeholder na tabela.
- Apresente resultado em português do Brasil, com valores em R$ e datas em dd/mm/yyyy.

""",
            input=prompt,
            tools=[
                {
                    "type": "web_search"
                }
            ],
            tool_choice="auto",
        )

        if not resposta.output_text:
            return "A API não retornou resultado."

        return resposta.output_text

    except Exception as erro:
        return f"Erro ao consultar GPT: {erro}"

@st.dialog("NÚMERO GCPJ É OBRIGATÓRIO")
def dialog_gcpj_obrigatorio():

    if st.button("OK", key="btn_ok_dialog_gcpj_obrigatorio"):
        st.rerun()

def validar_usuario_apps_script(usuario, senha):
    url = st.secrets.get("LEXCALX_HISTORICO_URL", "")
    token = st.secrets.get("LEXCALX_HISTORICO_TOKEN", "")

    if not url or not token:
        st.error("Configuração do login não encontrada nos secrets.")
        return False

    payload = {
        "acao": "validar_login",
        "token": token,
        "usuario": usuario,
        "senha": senha,
    }

    try:
        resposta = requests.post(url, json=payload, timeout=20)
        resposta.raise_for_status()

        retorno = resposta.json()

        if not retorno.get("ok"):
            st.error(f"Erro ao validar login: {retorno.get('erro', 'Erro desconhecido')}")
            return False

        return bool(retorno.get("autorizado"))

    except Exception as erro:
        st.error(f"Erro ao validar login: {erro}")
        return False

# =========================================================
# INTERFACE
# =========================================================
st.markdown(
    """
    <div class="hero-title-box">
        <div class="main-title">LexCalx</div>
        <div class="subtitle">
            Cálculos jurídicos
            <span class="subtitle-helper">?</span>
            <span class="subtitle-tooltip">
                LEXCALX realiza análise de decisões judiciais líquidáveis ou passíveis de cálculos.<br>
                Para maior precisão, preencha assertivamente o máximo possível de campos de data.<br>
                Me envia o arquivo da decisão em pdf nativo ou cole no campo final o texto da decisão.
            </span>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# =========================================================
# LOGIN
# =========================================================
if "lexcalx_logado" not in st.session_state:
    st.session_state.lexcalx_logado = False

if "usuario_logado" not in st.session_state:
    st.session_state.usuario_logado = ""

if not st.session_state.lexcalx_logado:
    col_login_esq, col_login_centro, col_login_dir = st.columns([1.2, 1, 1.2])

    with col_login_centro:
        usuario = st.text_input(
            "Login:",
            key="login_usuario",
        )

        senha = st.text_input(
            "Senha:",
            type="password",
            key="login_senha",
        )

        entrar = st.button(
            "Entrar",
            key="btn_entrar_login",
        )

        if entrar:
            if validar_usuario_apps_script(usuario, senha):
                st.session_state.lexcalx_logado = True
                st.session_state.usuario_logado = usuario.strip()
                st.rerun()
            else:
                st.error("Login ou senha inválidos.")

    st.stop()

col_gcpj, col_processo = st.columns(2)

with col_gcpj:
    gcpj = st.text_input(
        "GCPJ: (campo obrigatório)",
        key="input_gcpj",
        help=(
            "Campo obrigatório.Esse dado será usado para controle interno e constará apenas no histórico completo, "
        ),
    )

with col_processo:
    numero_processo = st.text_input(
        "Número do processo:",
        key="input_numero_processo",
    )

col1, col2, col3 = st.columns(3)

with col1:
    data_distribuicao = st.date_input(
        "Data da distribuição",
        value=None,
        format="DD/MM/YYYY",
    )

    data_citacao = st.date_input(
        "Data da citação",
        value=None,
        format="DD/MM/YYYY",
    )

    data_evento_danoso = st.date_input(
        "Data do evento danoso",
        value=None,
        format="DD/MM/YYYY",
    )

with col2:
    data_juros_dano_moral = st.date_input(
        "Data juros dano moral",
        value=None,
        format="DD/MM/YYYY",
    )

    data_correcao_dano_moral = st.date_input(
        "Data correção monetária dano moral",
        value=None,
        format="DD/MM/YYYY",
    )

    data_juros_dano_material = st.date_input(
        "Data juros dano material",
        value=None,
        format="DD/MM/YYYY",
    )

with col3:
    data_correcao_dano_material = st.date_input(
        "Data correção monetária dano material",
        value=None,
        format="DD/MM/YYYY",
    )

    data_decisao = st.date_input(
        "Data da decisão",
        value=None,
        format="DD/MM/YYYY",
    )

    data_publicacao = st.date_input(
        "Data da publicação",
        value=None,
        format="DD/MM/YYYY",
    )

    # =========================================================
# DANO MATERIAL
# =========================================================
if "qtd_linhas_dano_material" not in st.session_state:
    st.session_state.qtd_linhas_dano_material = 1


st.markdown("### Dano material")

linhas_dano_material = []

for i in range(st.session_state.qtd_linhas_dano_material):
    col_dmt1, col_dmt2, col_dmt3 = st.columns([1, 0.8, 1])

    with col_dmt1:
        valor_dano_material = st.text_input(
            f"Valor do dano material {i + 1}",
            placeholder="Ex.: 5.000,00",
            key=f"valor_dano_material_{i}",
            on_change=normalizar_campo_valor_dano_material,
            args=(i,),
    )

    with col_dmt2:
        parcelas_dano_material = st.number_input(
            f"Parcelas {i + 1}",
            min_value=1,
            value=None,
            step=1,
            key=f"parcelas_dano_material_{i}",
        )

    with col_dmt3:
        data_inicial_dano_material = st.date_input(
            f"Data inicial {i + 1}",
            value=None,
            format="DD/MM/YYYY",
            key=f"data_inicial_dano_material_{i}",
        )

    linhas_dano_material.append(
        {
            "valor": valor_dano_material,
            "parcelas": parcelas_dano_material,
            "data_inicial": data_inicial_dano_material,
        }
    )


col_btn_add_dm, col_btn_remove_dm = st.columns(2)

with col_btn_add_dm:
    if st.button(
        "Adicionar outro dano material",
        key="btn_add_dano_material",
    ):
        st.session_state.qtd_linhas_dano_material += 1
        st.rerun()

with col_btn_remove_dm:
    if st.session_state.qtd_linhas_dano_material > 1:
        if st.button(
            "Remover último dano material",
            key="btn_remove_dano_material",
        ):
            idx_remover = st.session_state.qtd_linhas_dano_material - 1

            st.session_state.pop(f"valor_dano_material_{idx_remover}", None)
            st.session_state.pop(f"parcelas_dano_material_{idx_remover}", None)
            st.session_state.pop(f"data_inicial_dano_material_{idx_remover}", None)

            st.session_state.qtd_linhas_dano_material -= 1
            st.rerun()
    else:
        st.button(
            "Remover último dano material",
            key="btn_remove_dano_material_disabled",
            disabled=True,
        )

st.markdown("### Configurações do cálculo")

col_config1, col_config2 = st.columns(2)

with col_config1:
    inserir_art_523 = st.selectbox(
        "Inserir art. 523 (multa 10% e honorários 10%)?",
        options=["Não", "Multa 10%", "Honorários 10%", "Ambos"],
        index=0,
        key="inserir_art_523",
    )

with col_config2:
    data_calculo_usuario = st.date_input(
        "Atualizar o cálculo até:",
        value=obter_data_calculo(),
        format="DD/MM/YYYY",
        key="data_calculo_usuario",
    )

st.divider()

arquivos_pdf = st.file_uploader(
    "Me envia o PDF NATIVO da decisão:(se o PDF for escaneado, imagem, ou tiver texto não extraível, o resultado vem vazio.)",
    type=["pdf"],
    accept_multiple_files=True,
)

texto_colado = st.text_area(
    "Ou cole aqui o texto da decisão:",
    height=260,
)

if "resultado_lexcalx" not in st.session_state:
    st.session_state.resultado_lexcalx = None
    st.session_state.resumo_lexcalx = None
    st.session_state.historico_lexcalx = None
    st.session_state.dados_lexcalx = None
    st.session_state.planilha_lexcalx = None
    st.session_state.pdf_resumo_lexcalx = None
    st.session_state.pdf_lexcalx = None
    st.session_state.nome_planilha_lexcalx = "historico_lexcalx.xlsx"
    st.session_state.nome_pdf_resumo_lexcalx = "calculo_resumo_processo.pdf"
    st.session_state.nome_pdf_lexcalx = "historico_completo_lexcalx.pdf"

executar = st.button("Executar cálculo")

if executar:
    st.session_state.resultado_lexcalx = None
    st.session_state.resumo_lexcalx = None
    st.session_state.historico_lexcalx = None
    st.session_state.dados_lexcalx = None
    st.session_state.planilha_lexcalx = None
    st.session_state.pdf_resumo_lexcalx = None
    st.session_state.pdf_lexcalx = None

    if not gcpj.strip():
        dialog_gcpj_obrigatorio()
        st.stop()

    textos_pdfs = []

    for idx, arquivo_pdf in enumerate(arquivos_pdf, start=1):
        texto_extraido = extrair_texto_pdf(arquivo_pdf)

        if texto_extraido:
            textos_pdfs.append(
                f"\n\n[INÍCIO DO PDF {idx} - {arquivo_pdf.name}]\n"
                f"{texto_extraido}\n"
                f"[FIM DO PDF {idx} - {arquivo_pdf.name}]\n"
            )

    texto_extraido_pdf = "\n\n".join(textos_pdfs).strip()

    textos_decisao = []

    if texto_extraido_pdf:
        textos_decisao.append(texto_extraido_pdf)

    if texto_colado.strip():
        textos_decisao.append(
            "\n\n[INÍCIO DO TEXTO COLADO PELO USUÁRIO]\n"
            f"{texto_colado.strip()}\n"
            "[FIM DO TEXTO COLADO PELO USUÁRIO]\n"
        )

    texto_decisao = "\n\n".join(textos_decisao).strip()

    if not texto_decisao:
        st.warning("Vincule ao menos um PDF ou cole o texto da decisão antes de executar o cálculo.")
        st.stop()

    danos_materiais_formatados = []

    for idx, linha in enumerate(linhas_dano_material, start=1):
        valor = formatar_valor_texto_br(linha["valor"])
        parcelas = linha["parcelas"]
        data_inicial = linha["data_inicial"]

        if valor is not None or parcelas is not None or data_inicial is not None:
            danos_materiais_formatados.append(
                f"{idx}. Valor: R$ {valor} | "
                f"Parcelas: {parcelas if parcelas is not None else ''} | "
                f"Data inicial: {formatar_data_br(data_inicial)}"
            )

    danos_materiais_texto = (
        "\n".join(danos_materiais_formatados)
        if danos_materiais_formatados
        else "Não informado."
    )

    data_calculo = data_calculo_usuario

    dados = {
        "gcpj": gcpj.strip(),
        "numero_processo": numero_processo.strip(),
        "data_calculo": formatar_data_br(data_calculo),
        "inserir_art_523": inserir_art_523,
        "percentual_multa_art_523": "10%" if inserir_art_523 in ["Multa 10%", "Ambos"] else "Não aplicável",
        "percentual_honorarios_art_523": "10%" if inserir_art_523 in ["Honorários 10%", "Ambos"] else "Não aplicável",
        "percentual_total_art_523": (
            "20%" if inserir_art_523 == "Ambos"
            else "10%" if inserir_art_523 in ["Multa 10%", "Honorários 10%"]
            else "Não aplicável"
        ),
        "data_distribuicao": formatar_data_br(data_distribuicao),
        "data_citacao": formatar_data_br(data_citacao),
        "data_evento_danoso": formatar_data_br(data_evento_danoso),
        "data_juros_dano_moral": formatar_data_br(data_juros_dano_moral),
        "data_correcao_dano_moral": formatar_data_br(data_correcao_dano_moral),
        "data_juros_dano_material": formatar_data_br(data_juros_dano_material),
        "data_correcao_dano_material": formatar_data_br(data_correcao_dano_material),
        "data_decisao": formatar_data_br(data_decisao),
        "data_publicacao": formatar_data_br(data_publicacao),
        "danos_materiais": danos_materiais_texto,
    }

    registrar_solicitacao_calculo(
        usuario_logado=st.session_state.usuario_logado,
        gcpj=gcpj.strip(),
        numero_processo=numero_processo.strip(),
        data_calculo=formatar_data_br(data_calculo),
    )

    prompt = montar_prompt(dados, texto_decisao)

    with st.spinner("Realizando o cálculo. Por favor, aguarde."):
        resultado_gpt = chamar_gpt(prompt)
        resumo_processo, historico_completo = separar_resultado_lexcalx(resultado_gpt)

    nome_planilha = "historico_completo_lexcalx.xlsx"
    nome_pdf = "historico_completo_lexcalx.pdf"
    nome_pdf_resumo = "calculo_resumo_processo.pdf"

    if numero_processo:
        numero_limpo = (
            numero_processo
            .replace(".", "")
            .replace("-", "")
            .replace("/", "")
            .replace("\\", "")
            .replace(" ", "_")
        )

        nome_planilha = f"historico_completo_lexcalx_{numero_limpo}.xlsx"
        nome_pdf = f"historico_completo_lexcalx_{numero_limpo}.pdf"
        nome_pdf_resumo = f"calculo_resumo_processo_{numero_limpo}.pdf"

    planilha_calculo = gerar_planilha_calculo(
        dados=dados,
        resultado_gpt=historico_completo,
    )

    pdf_resumo = gerar_pdf_resumo_processo(
        dados=dados,
        resumo_processo=resumo_processo,
    )

    pdf_calculo = gerar_pdf_calculo(
        dados=dados,
        resultado_gpt=historico_completo,
    )

    st.session_state.resultado_lexcalx = resultado_gpt
    st.session_state.resumo_lexcalx = resumo_processo
    st.session_state.historico_lexcalx = historico_completo
    st.session_state.dados_lexcalx = dados
    st.session_state.planilha_lexcalx = planilha_calculo.getvalue()
    st.session_state.pdf_resumo_lexcalx = pdf_resumo.getvalue()
    st.session_state.pdf_lexcalx = pdf_calculo.getvalue()
    st.session_state.nome_planilha_lexcalx = nome_planilha
    st.session_state.nome_pdf_resumo_lexcalx = nome_pdf_resumo
    st.session_state.nome_pdf_lexcalx = nome_pdf

if st.session_state.resumo_lexcalx:
    st.markdown("### Versão resumida para juntada no processo")

    resumo_html = markdown_resumo_para_html(st.session_state.resumo_lexcalx)

    st.markdown(
        f'<div class="resumo-processo-box">{resumo_html}</div>',
        unsafe_allow_html=True,
    )

    col_down1, col_down2, col_down3 = st.columns(3)

    with col_down1:
        st.download_button(
            label="Baixar PDF para juntar aos autos",
            data=st.session_state.pdf_resumo_lexcalx,
            file_name=st.session_state.nome_pdf_resumo_lexcalx,
            mime="application/pdf",
            key="btn_baixar_pdf_resumo_processo",
        )

    with col_down2:
        st.download_button(
            label="Baixar histórico completo PDF",
            data=st.session_state.pdf_lexcalx,
            file_name=st.session_state.nome_pdf_lexcalx,
            mime="application/pdf",
            key="btn_baixar_pdf_historico_completo",
        )

    with col_down3:
        st.download_button(
            label="Baixar histórico Excel",
            data=st.session_state.planilha_lexcalx,
            file_name=st.session_state.nome_planilha_lexcalx,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="btn_baixar_planilha_historico_completo",
        )

    with st.expander("Ver histórico completo do cálculo"):
        historico_html = markdown_resumo_para_html(st.session_state.historico_lexcalx)

        st.markdown(
            f'<div class="resumo-processo-box">{historico_html}</div>',
            unsafe_allow_html=True,
        )